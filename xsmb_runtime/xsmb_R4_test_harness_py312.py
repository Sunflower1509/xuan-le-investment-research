#!/usr/bin/env python3
from __future__ import annotations
import copy, hashlib, importlib.util, json, os, shutil, subprocess, sys, tempfile, time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable
from openpyxl import load_workbook

RUNNER=Path('/mnt/data/xsmb_reference_runner_R4_snapshot.py')
SOURCE_WB=Path('/mnt/data/XSMB_RECONCILIATION_R5_1200_v2_1_R3_HARDENED.xlsx')
REAL_EVIDENCE=Path('/mnt/data/settlement_evidence_20260920_4source.json')
OUT=Path('/mnt/data/xsmb_R4_test_output')
OUT.mkdir(exist_ok=True)

spec=importlib.util.spec_from_file_location('xsmb_r4',RUNNER)
r=importlib.util.module_from_spec(spec); sys.modules['xsmb_r4']=r; spec.loader.exec_module(r)


def sha(path: Path)->str:
    return r.sha256_file(path) if path.exists() else 'MISSING'


def set_activation(wb,key,val):
    sh=wb['MODEL_ACTIVATION']
    for rr in range(2,sh.max_row+1):
        if str(sh.cell(rr,1).value or '')==key:
            sh.cell(rr,2).value=val; return
    rr=sh.max_row+1; sh.cell(rr,1).value=key; sh.cell(rr,2).value=val


def make_base_fixture(dest: Path)->Path:
    shutil.copy2(SOURCE_WB,dest)
    wb=load_workbook(dest)
    sh=wb['RECONCILIATION_1200']
    # Remove appended row 1201. Keep immutable MASTER values/hashes exactly as supplied.
    for c in range(1,14): sh.cell(1202,c).value=None
    # Production canonical file in File Library is 1200/1200 reconciled. For SOFTWARE TEST FIXTURE ONLY,
    # clear stale mirror flags so runner mechanics can be tested without certifying this workbook as production data.
    for rr in range(2,1202):
        if sh.cell(rr,1).value in (None,''): continue
        sh.cell(rr,11).value=0; sh.cell(rr,12).value=0
        if str(sh.cell(rr,10).value or '')=='PENDING_MACHINE_COMPARE':
            sh.cell(rr,7).value='TEST_FIXTURE_CANONICAL_R5_VERIFIED'
            sh.cell(rr,8).value='Canonical File Library R5 assertion (software-test fixture only)'
            sh.cell(rr,9).value='file_000000003afc81f7abe5e69f353080c2'
            sh.cell(rr,10).value='MATCH_TEST_FIXTURE'
            sh.cell(rr,13).value='SOFTWARE TEST FIXTURE ONLY; not production provenance.'
    led=wb['PROSPECTIVE_LEDGER']
    for rr in range(2,led.max_row+1):
        if led.cell(rr,1).value not in (None,''):
            for c in range(1,17): led.cell(rr,c).value=None
    set_activation(wb,'reference_runner_sha256',r.runner_sha256())
    set_activation(wb,'implementation_gate',r.IMPLEMENTATION_GATE_PASS)
    set_activation(wb,'first_forecast_status','NOT_STARTED')
    set_activation(wb,'reproducibility_gate','TEST_FIXTURE_R3')
    wb.save(dest); wb.close(); return dest


def make_test_evidence(path: Path)->Path:
    tests=[{'id':tid,'status':'PASS'} for tid in r.ALL_TEST_IDS]
    payload={'artifact_type':'TEST_EVIDENCE_MANIFEST_V1','test_suite_revision':r.TEST_SUITE_REVISION,'spec_hash':r.SPEC_HASH,'runner_sha256':r.runner_sha256(),'master_content_sha256':r.FROZEN_MASTER_1200_SHA256,'runtime_exact_production_match':True,'passed':34,'failed':0,'tests':tests,'testing_fixture':True}
    obj=r._manifest_with_hash(payload); path.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8'); return path


def add_ledger_row(path: Path, vals: list[Any]):
    wb=load_workbook(path); sh=wb['PROSPECTIVE_LEDGER']; rr=2
    while sh.cell(rr,1).value not in (None,''): rr+=1
    for c,v in enumerate(vals,1): sh.cell(rr,c).value=v
    wb.save(path); wb.close()


def seed_forecast(path: Path, manifest_dir: Path, forecast_date='2026-09-20', data_cutoff='2026-09-19')->dict[str,Any]:
    state=r.load_workbook_state(path,manifest_dir)
    dh=r.data_hash(state.draws); fid=f"XSMB-{forecast_date.replace('-','')}-v2.1-LIVE-001-{dh[:12]}-{r.SPEC_HASH[:12]}"
    p={f'{i:02d}':float(0.05 + i/10000) for i in range(100)}
    top5=['99','98','97','96','95']; gen='2026-09-20T17:00:00+07:00'
    vals=[fid,forecast_date,gen,data_cutoff,dh,'v2.1-LIVE-001','NO VERIFIED EDGE','M7_logistic_l2',r.canonical_json(p),r.canonical_json(top5),None,None,None,None,None,'LOCKED_PRE_DRAW']
    add_ledger_row(path,vals)
    payload={'artifact_type':'XSMB_FORECAST_MANIFEST_V1','forecast_id':fid,'forecast_date':forecast_date,'generated_at_local':gen,'data_cutoff':data_cutoff,'data_hash':dh,'model_spec_version':'v2.1-LIVE-001','research_state':'NO VERIFIED EDGE','model':'M7_logistic_l2','probabilities':p,'top5_ranked':top5,'m0_uniform':0.24,'record_status':'LOCKED_PRE_DRAW','runner_sha256':r.runner_sha256(),'spec_hash':r.SPEC_HASH}
    manifest_dir.mkdir(parents=True,exist_ok=True); mp=manifest_dir/f'{fid}.forecast.json'; mp.write_text(json.dumps(r._manifest_with_hash(payload),ensure_ascii=False,indent=2),encoding='utf-8')
    return {'forecast_id':fid,'data_hash':dh,'manifest':mp,'probabilities':p,'top5':top5}


def bad_evidence(path: Path, mode: str):
    good=json.loads(REAL_EVIDENCE.read_text())
    if mode=='two_mismatch':
        src=copy.deepcopy(good['sources'][:2]); src[1]['lotto_27']=src[1]['lotto_27'].copy(); src[1]['lotto_27'][0]='00'; src[1]['prizes_27']=None; src[1]['canonical_order']=False
    elif mode=='same_family':
        src=copy.deepcopy(good['sources'][:2]); src[1]['source_family']=src[0]['source_family']
    elif mode=='tie3':
        src=copy.deepcopy(good['sources'][:3])
        for i,s in enumerate(src):
            s['lotto_27']=s['lotto_27'].copy(); s['lotto_27'][0]=f'{i:02d}'; s['prizes_27']=None; s['canonical_order']=False
    elif mode=='conflict_rerun':
        src=copy.deepcopy(good['sources'])
        for s in src:
            s['lotto_27']=s['lotto_27'].copy(); s['lotto_27'][0]='00'; s['prizes_27']=None; s['canonical_order']=False
    else: raise ValueError(mode)
    # Keep altered evidence internally self-consistent so tests reach the intended consensus/idempotency gate.
    for item in src:
        item['extract_sha256']=r.sha256_bytes(r.canonical_json({'prizes_27':item.get('prizes_27'),'lotto_27':item['lotto_27']}).encode('utf-8'))
    path.write_text(json.dumps({'actual_date':'2026-09-20','sources':src},ensure_ascii=False,indent=2),encoding='utf-8')
    return path


def fake_prediction(draws,target):
    p={f'{i:02d}':float((i+1)/1000) for i in range(100)}
    return {'probabilities':p,'top5':['99','98','97','96','95'],'probability_sha256':r.probability_sha256(list(p.values())),'fit_target_range':[61,target-201],'calibration_target_range':[target-200,target-1],'m0_uniform':0.24}


def helper_script(path: Path):
    path.write_text(r'''import importlib.util,sys,json,os
from pathlib import Path
runner=Path(sys.argv[1]); spec=importlib.util.spec_from_file_location("rr",runner); rr=importlib.util.module_from_spec(spec); sys.modules["rr"]=rr; spec.loader.exec_module(rr)
rr.assert_runtime=lambda: rr.runtime_manifest()
rr._validate_test_evidence=lambda *a,**k: {"payload":{"testing":True}}
def fp(draws,target):
 p={f"{i:02d}":float((i+1)/1000) for i in range(100)}
 return {"probabilities":p,"top5":["99","98","97","96","95"],"probability_sha256":rr.probability_sha256(list(p.values())),"fit_target_range":[61,target-201],"calibration_target_range":[target-200,target-1],"m0_uniform":0.24}
rr.fit_predict=fp; rr.time_series_validation=lambda d,t:{"mode":"TEST_FAST"}
cmd=sys.argv[2]
try:
 if cmd=="run": out=rr.run_forecast(sys.argv[3],"2026-09-21","2026-09-20",sys.argv[4],sys.argv[5])
 elif cmd=="settle": out=rr.settle_forecast(sys.argv[3],sys.argv[6],"2026-09-20",sys.argv[7],sys.argv[4])
 else: raise RuntimeError(cmd)
 print(json.dumps(out))
except Exception as e:
 if hasattr(e,"payload"): print(json.dumps(e.payload()),file=sys.stderr)
 else: print(repr(e),file=sys.stderr)
 sys.exit(2)
''',encoding='utf-8')


def record(test_id, name, expected, fn: Callable[[],Any], input_path: Path|None=None, manifest_paths: list[Path]|None=None):
    before=sha(input_path) if input_path else None
    m_before={str(p):sha(p) for p in (manifest_paths or [])}
    status='FAIL'; observed=None; error=None
    try:
        observed=fn(); status='PASS'
    except AssertionError as e:
        error=f'AssertionError: {e}'
    except Exception as e:
        error=f'{type(e).__name__}: {e}'
    after=sha(input_path) if input_path else None
    m_after={str(p):sha(p) for p in (manifest_paths or [])}
    result={'id':test_id,'name':name,'expected':expected,'observed':observed,'error':error,'input_hash_before':before,'input_hash_after':after,'manifest_hashes_before':m_before,'manifest_hashes_after':m_after,'status':status}
    RESULTS.append(result); print(test_id,status, error or '')
    return status=='PASS'

RESULTS=[]
ROOT=Path(tempfile.mkdtemp(prefix='xsmb-r3-tests-',dir='/mnt/data'))
BASE=make_base_fixture(ROOT/'base.xlsx')
MAN=ROOT/'manifests'; MAN.mkdir()
TE=make_test_evidence(ROOT/'test_fixture_evidence.json')

# Sanity: test fixture preserves frozen master content hash.
state=r.load_workbook_state(BASE,MAN)
assert r.canonical_master_hash(state.draws)==r.FROZEN_MASTER_1200_SHA256 and len(state.draws.dates)==1200

# R1 real M7 same process.
def tR1():
 s=r.load_workbook_state(BASE,MAN); a=r.fit_predict(s.draws,1201); b=r.fit_predict(s.draws,1201)
 va=list(a['probabilities'].values()); vb=list(b['probabilities'].values())
 assert va==vb and a['top5']==b['top5']
 return {'probability_sha256':a['probability_sha256'],'top5':a['top5']}
record('R1','same immutable input, same process','P bitwise equal; TOP5 equal',tR1,BASE)

# R2 real M7 in two fresh helper processes without production-runtime assertion; this tests deterministic math only.
def tR2():
 helper=ROOT/'predict_helper.py'; helper.write_text(r'''import importlib.util,sys,json
p=sys.argv[1]; spec=importlib.util.spec_from_file_location("m",p); m=importlib.util.module_from_spec(spec); sys.modules["m"]=m; spec.loader.exec_module(m)
s=m.load_workbook_state(sys.argv[2],sys.argv[3]); pred=m.fit_predict(s.draws,1201); print(json.dumps({"sha":pred["probability_sha256"],"top5":pred["top5"]}))''')
 cmd=[sys.executable,str(helper),str(RUNNER),str(BASE),str(MAN)]
 a=json.loads(subprocess.run(cmd,check=True,capture_output=True,text=True).stdout); b=json.loads(subprocess.run(cmd,check=True,capture_output=True,text=True).stdout)
 assert a==b
 return a
record('R2','same immutable input, fresh process','probability SHA and TOP5 equal',tR2,BASE)

# R3 simulate a runtime mismatch independent of the host runtime.
def tR3():
 old=r.RUNTIME_LOCK['python_major_minor']; r.RUNTIME_LOCK['python_major_minor']='0.0'
 try:
  try:r.assert_runtime(); raise AssertionError('runtime mismatch unexpectedly passed')
  except r.RunnerError as e:
   assert 'runtime lock mismatch' in e.reason and any('python=' in x for x in e.details)
   return e.payload()
 finally:r.RUNTIME_LOCK['python_major_minor']=old
record('R3','runtime/package mismatch','hard fail, no Ledger write',tR3,BASE)

# R4 import-time thread lock violation.
def tR4():
 env=os.environ.copy(); env['OMP_NUM_THREADS']='2'
 p=subprocess.run([sys.executable,str(RUNNER),'print-spec'],env=env,capture_output=True,text=True)
 assert p.returncode!=0 and 'thread environment was not locked' in (p.stderr+p.stdout)
 return {'returncode':p.returncode,'message_detected':True}
record('R4','thread env != 1','hard fail before numerical import/fit',tR4,BASE)

# R5 SPEC hash mismatch.
def tR5():
 old=r.EXPECTED_SPEC_HASH; r.EXPECTED_SPEC_HASH='deadbeef'
 try:
  try:r.preflight_forecast(BASE,'2026-09-20','2026-09-19',MAN,TE,now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); raise AssertionError('passed')
  except r.RunnerError as e: assert e.status.startswith('HASH MISMATCH'); return e.payload()
 finally:r.EXPECTED_SPEC_HASH=old
record('R5','spec_hash mismatch','hard fail, no fit',tR5,BASE)

# R6 duplicate identity with wrong data_hash.
def tR6():
 p=ROOT/'r6.xlsx'; shutil.copy2(BASE,p); s=r.load_workbook_state(p,MAN); dh=r.data_hash(s.draws); fid=f"XSMB-20260920-v2.1-LIVE-001-{dh[:12]}-{r.SPEC_HASH[:12]}"
 add_ledger_row(p,[fid,'2026-09-20','2026-09-20T17:00:00+07:00','2026-09-19','0'*64,'v2.1-LIVE-001','NO VERIFIED EDGE','M7_logistic_l2','{}','[]',None,None,None,None,None,'LOCKED_PRE_DRAW'])
 try:r.preflight_forecast(p,'2026-09-20','2026-09-19',MAN,TE,now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); raise AssertionError('passed')
 except r.RunnerError as e: assert 'duplicate target date' in e.reason; return e.payload()
record('R6','data_hash mismatch','identity conflict, no fit',tR6)

# R7 target result already present.
def tR7():
 p=ROOT/'r7.xlsx'; shutil.copy2(BASE,p); md=ROOT/'r7m'; md.mkdir(); r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md)
 # activation runner hash was preserved, implementation gate pass.
 try:r.preflight_forecast(p,'2026-09-20','2026-09-19',md,TE,now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); raise AssertionError('passed')
 except r.RunnerError as e: assert e.status.startswith('TARGET LEAKAGE'); return e.payload()
record('R7','target result present','leakage fail, no forecast',tR7)

# R8 cutoff.
def tR8():
 try:r.preflight_forecast(BASE,'2026-09-20','2026-09-19',MAN,TE,now=datetime.fromisoformat('2026-09-20T18:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); raise AssertionError('passed')
 except r.RunnerError as e: assert 'cutoff' in e.reason; return e.payload()
record('R8','current time >= internal cutoff','no forecast',tR8,BASE)

# R9 duplicate identical target -> idempotent ALREADY_EXISTS.
def tR9():
 p=ROOT/'r9.xlsx'; shutil.copy2(BASE,p); s=r.load_workbook_state(p,MAN); dh=r.data_hash(s.draws); fid=f"XSMB-20260920-v2.1-LIVE-001-{dh[:12]}-{r.SPEC_HASH[:12]}"
 add_ledger_row(p,[fid,'2026-09-20','2026-09-20T17:00:00+07:00','2026-09-19',dh,'v2.1-LIVE-001','NO VERIFIED EDGE','M7_logistic_l2','{}','[]',None,None,None,None,None,'LOCKED_PRE_DRAW'])
 out=r.preflight_forecast(p,'2026-09-20','2026-09-19',MAN,TE,now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); assert out['status']=='ALREADY_EXISTS'; return out
record('R9','duplicate target/date','idempotent read or conflict',tR9)

# R10 prior forecast unsettled.
def tR10():
 p=ROOT/'r10.xlsx'; shutil.copy2(BASE,p)
 add_ledger_row(p,['OLD','2026-09-19','2026-09-19T17:00:00+07:00','2026-09-18','x','v2.1-LIVE-001','NO VERIFIED EDGE','M7_logistic_l2','{}','[]',None,None,None,None,None,'LOCKED_PRE_DRAW'])
 try:r.preflight_forecast(p,'2026-09-20','2026-09-19',MAN,TE,now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); raise AssertionError('passed')
 except r.RunnerError as e: assert 'DATA CONTINUITY FAIL' in e.reason; return e.payload()
record('R10','missing prior draw/unsettled state','no forecast',tR10)

# R11 non-finite prediction rejected.
def tR11():
 class S:
  def transform(self,x): return x
 class B:
  def decision_function(self,x):
   import numpy as np
   return np.zeros(len(x), dtype=float)
 class C:
  def predict_proba(self,x):
   import numpy as np
   a=np.empty((len(x),2)); a[:,0]=0.5; a[:,1]=np.nan; return a
 s=r.load_workbook_state(BASE,MAN)
 try:r._predict_components(s.draws,1201,S(),B(),C()); raise AssertionError('passed')
 except r.RunnerError as e: assert 'invalid probability vector' in e.reason; return e.payload()
record('R11','convergence/non-finite output','no forecast',tR11,BASE)

# R12 missing test evidence.
def tR12():
 try:r.preflight_forecast(BASE,'2026-09-20','2026-09-19',MAN,ROOT/'missing.json',now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=False); raise AssertionError('passed')
 except r.RunnerError as e: assert 'missing manifest' in e.reason; return e.payload()
record('R12','missing live implementation/validation evidence','implementation gate fail',tR12,BASE)

# R13 settle no row.
def tR13():
 try:r.settle_forecast(BASE,'NO-SUCH-ID','2026-09-20',REAL_EVIDENCE,MAN); raise AssertionError('passed')
 except r.RunnerError as e: assert e.status.startswith('NO PRE-DRAW RECORD'); return e.payload()
record('R13','SETTLE without pre-draw row','backfill forbidden',tR13,BASE)

# R14 two-source mismatch, no third -> settlement blocked and no workbook write.
def tR14():
 p=ROOT/'r14.xlsx'; shutil.copy2(BASE,p); md=ROOT/'r14m'; md.mkdir(); f=seed_forecast(p,md); ev=bad_evidence(ROOT/'r14e.json','two_mismatch'); before=sha(p)
 try:r.settle_forecast(p,f['forecast_id'],'2026-09-20',ev,md); raise AssertionError('passed')
 except r.RunnerError as e: assert e.status=='SETTLEMENT_BLOCKED' and sha(p)==before; return e.payload()
record('R14','PDF/source multiset mismatch without third source','settlement blocked',tR14)

# R15 mutate P/TOP5/hash -> reject settlement.
def tR15():
 p=ROOT/'r15.xlsx'; shutil.copy2(BASE,p); md=ROOT/'r15m'; md.mkdir(); f=seed_forecast(p,md); wb=load_workbook(p); sh=wb['PROSPECTIVE_LEDGER']; sh.cell(2,9).value='{}'; wb.save(p); wb.close(); before=sha(p)
 try:r.settle_forecast(p,f['forecast_id'],'2026-09-20',REAL_EVIDENCE,md); raise AssertionError('passed')
 except r.RunnerError as e: assert 'immutable forecast fields' in e.reason and sha(p)==before; return e.payload()
record('R15','settlement tries to alter P/TOP5/hash','reject transaction',tR15)

# C01 change historical lotto + repair row hash so frozen MASTER hash catches it.
def tC01():
 p=ROOT/'c01.xlsx'; shutil.copy2(BASE,p); wb=load_workbook(p); sh=wb['RECONCILIATION_1200']; rr=101; idx=int(sh.cell(rr,1).value); d=str(sh.cell(rr,2).value); vals=str(sh.cell(rr,4).value).split(); vals[0]='00' if vals[0] != '00' else '01'; sh.cell(rr,4).value=' '.join(vals); sh.cell(rr,5).value=r.sha256_bytes(r.canonical_json({'draw_index':idx,'date':d,'lotto_27':vals}).encode()); sh.cell(rr,6).value=r.sha256_bytes(r.canonical_json(sorted(vals)).encode()); wb.save(p); wb.close()
 try:r.load_workbook_state(p,MAN); raise AssertionError('passed')
 except r.RunnerError as e: assert 'MASTER 1,200 fingerprint mismatch' in e.reason; return e.payload()
record('C01','mutate historical lotto','MASTER HASH FAIL',tC01)

# C02 falsify R5.
def tC02():
 p=ROOT/'c02.xlsx'; shutil.copy2(BASE,p); wb=load_workbook(p); sh=wb['R5_GATE'];
 for rr in range(1,sh.max_row+1):
  if str(sh.cell(rr,1).value or '')=='R5 FINAL': sh.cell(rr,4).value='FAIL'
 wb.save(p); wb.close()
 try:r.load_workbook_state(p,MAN); raise AssertionError('passed')
 except r.RunnerError as e: assert 'R5 FINAL' in e.reason; return e.payload()
record('C02','falsify R5 gate','hard fail',tC02)

# C03 invalid first forecast status.
def tC03():
 p=ROOT/'c03.xlsx'; shutil.copy2(BASE,p); wb=load_workbook(p); set_activation(wb,'first_forecast_status','BROKEN'); wb.save(p); wb.close()
 try:r.load_workbook_state(p,MAN); raise AssertionError('passed')
 except r.RunnerError as e: assert 'first_forecast_status' in e.reason; return e.payload()
record('C03','invalid first_forecast_status','hard fail',tC03)

# C04 prior SETTLEMENT_BLOCKED is continuity gap.
def tC04():
 p=ROOT/'c04.xlsx'; shutil.copy2(BASE,p); add_ledger_row(p,['OLD','2026-09-19','x','2026-09-18','x','v2.1-LIVE-001','NO VERIFIED EDGE','M7_logistic_l2','{}','[]',None,None,None,None,None,'SETTLEMENT_BLOCKED'])
 try:r.preflight_forecast(p,'2026-09-20','2026-09-19',MAN,TE,now=datetime.fromisoformat('2026-09-20T17:00:00+07:00'),skip_runtime=True,skip_test_evidence=True); raise AssertionError('passed')
 except r.RunnerError as e: assert 'DATA CONTINUITY FAIL' in e.reason; return e.payload()
record('C04','prior SETTLEMENT_BLOCKED','DATA CONTINUITY FAIL',tC04)

# C05 same family twice.
def tC05():
 ev=bad_evidence(ROOT/'c05.json','same_family'); con,detail=r._read_evidence(ev,'2026-09-20'); assert con is None and 'independent source families' in detail['reason']; return detail['reason']
record('C05','same source_family twice','insufficient consensus',tC05)

# C06 successful settlement preserves immutable fields.
def tC06():
 p=ROOT/'c06.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c06m'; md.mkdir(); f=seed_forecast(p,md); wb=load_workbook(p); sh=wb['PROSPECTIVE_LEDGER']; before=[sh.cell(2,c).value for c in range(1,11)]; wb.close(); out=r.settle_forecast(p,f['forecast_id'],'2026-09-20',REAL_EVIDENCE,md); wb=load_workbook(p); sh=wb['PROSPECTIVE_LEDGER']; after=[sh.cell(2,c).value for c in range(1,11)]; rec=wb['RECONCILIATION_1200']; count=sum(1 for rr in range(2,rec.max_row+1) if rec.cell(rr,1).value not in (None,'')); wb.close(); assert before==after and count==1201 and out['status']=='SETTLED_VERIFIED'; return {'immutable_equal':True,'draw_count':count}
record('C06','successful settlement immutable fields','no mutation + one append',tC06)

# C07 crash after workbook replace -> recovery completes manifest.
def tC07():
 p=ROOT/'c07.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c07m'; md.mkdir(); state=r.load_workbook_state(p,MAN); wb=load_workbook(p); set_activation(wb,'c07_marker','COMMITTED_AFTER_RECOVERY'); manifest=r._manifest_with_hash({'artifact_type':'TEST_TXN_MANIFEST_V1','marker':'C07'}); mp=md/'c07.json'; os.environ['XSMB_FAULT_POINT']='AFTER_WORKBOOK_REPLACE'
 try:
  try:r.commit_workbook_and_manifest(p,wb,mp,manifest,state.workbook_sha256,'C07_TEST'); raise AssertionError('fault did not fire')
  except RuntimeError as e: assert 'FAULT_INJECTION' in str(e)
 finally:
  os.environ.pop('XSMB_FAULT_POINT',None); wb.close()
 actions=r.recover_transactions(p); assert mp.exists() and r._validate_payload_manifest(mp,'TEST_TXN_MANIFEST_V1') and any(a['action']=='COMPLETED_COMMIT' for a in actions)
 wb=load_workbook(p); assert any(str(wb['MODEL_ACTIVATION'].cell(rr,1).value or '')=='c07_marker' for rr in range(2,wb['MODEL_ACTIVATION'].max_row+1)); wb.close(); return {'actions':actions}
record('C07','interrupted commit','recovery to one consistent committed state',tC07)

# C08 run-forecast idempotency using fast test monkeypatches.
def tC08():
 p=ROOT/'c08.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c08m'; md.mkdir(); r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); te=make_test_evidence(ROOT/'c08te.json')
 olda,oldt,oldf=r.assert_runtime,r._validate_test_evidence,r.fit_predict; oldv=r.time_series_validation
 r.assert_runtime=lambda:r.runtime_manifest(); r._validate_test_evidence=lambda *a,**k:{'payload':{}}; r.fit_predict=fake_prediction; r.time_series_validation=lambda d,t:{'mode':'TEST_FAST'}
 try:
  a=r.run_forecast(p,'2026-09-21','2026-09-20',md,te); b=r.run_forecast(p,'2026-09-21','2026-09-20',md,te)
 finally:r.assert_runtime, r._validate_test_evidence, r.fit_predict, r.time_series_validation=olda,oldt,oldf,oldv
 wb=load_workbook(p); led=wb['PROSPECTIVE_LEDGER']; rows=[rr for rr in range(2,led.max_row+1) if led.cell(rr,1).value not in (None,'')]; wb.close(); assert len(rows)==1 and b['status']=='ALREADY_EXISTS'; return {'rows':len(rows),'rerun_status':b['status']}
record('C08','identical RUN_FORECAST rerun','idempotent; no duplicate Ledger row',tC08)

# C09 ingest observed when forecast exists.
def tC09():
 p=ROOT/'c09.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c09m'; md.mkdir(); seed_forecast(p,md)
 try:r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); raise AssertionError('passed')
 except r.RunnerError as e: assert 'forecast row exists' in e.reason; return e.payload()
record('C09','ingest observed when forecast row exists','reject',tC09)

# C10 future observed date.
def tC10():
 ev=json.loads(REAL_EVIDENCE.read_text()); ev['actual_date']='2026-09-21'; pth=ROOT/'c10e.json'; pth.write_text(json.dumps(ev))
 try:r.ingest_observed(BASE,'2026-09-21',pth,ROOT/'c10m'); raise AssertionError('passed')
 except r.RunnerError as e: assert 'future draw' in e.reason; return e.payload()
record('C10','ingest observed future date','reject',tC10,BASE)

# C11 non-next observed date with fake current date 22 Sep.
def tC11():
 ev=json.loads(REAL_EVIDENCE.read_text()); ev['actual_date']='2026-09-21'; pth=ROOT/'c11e.json'; pth.write_text(json.dumps(ev)); olddt=r.datetime
 class FakeDateTime(datetime):
  @classmethod
  def now(cls,tz=None): return cls(2026,9,22,12,0,0,tzinfo=tz)
 r.datetime=FakeDateTime
 try:
  try:r.ingest_observed(BASE,'2026-09-21',pth,ROOT/'c11m'); raise AssertionError('passed')
  except r.RunnerError as e: assert 'not next required historical date' in e.reason; return e.payload()
 finally:r.datetime=olddt
record('C11','ingest observed not next required draw','reject',tC11,BASE)

# C12 evidence disagreement (3-way tie) blocks.
def tC12():
 p=ROOT/'c12.xlsx'; shutil.copy2(BASE,p); ev=bad_evidence(ROOT/'c12e.json','tie3'); md=ROOT/'c12m'
 try:r.ingest_observed(p,'2026-09-20',ev,md); raise AssertionError('passed')
 except r.RunnerError as e: assert e.status=='SETTLEMENT_BLOCKED'; return e.payload()
record('C12','ingest observed evidence disagreement','BLOCK',tC12)

# C13 successful ingest invariants.
def tC13():
 p=ROOT/'c13.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c13m'; out=r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); s=r.load_workbook_state(p,md); wb=load_workbook(p); led=wb['PROSPECTIVE_LEDGER']; rows=[rr for rr in range(2,led.max_row+1) if led.cell(rr,1).value not in (None,'')]; wb.close(); man=r._validate_payload_manifest(r._find_manifest_for_observed(md,'2026-09-20'),'XSMB_OBSERVED_NO_FORECAST_V1')['payload']; assert len(s.draws.dates)==1201 and not rows and man['p00_p99_generated'] is False and man['top5_generated'] is False; return {'draws':1201,'ledger_forecast_rows':0,'status':out['status']}
record('C13','successful ingest observed','append history; no Ledger/P/TOP5',tC13)

# C14 identical ingest rerun.
def tC14():
 p=ROOT/'c14.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c14m'; r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); out=r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); assert out['status']=='ALREADY_INGESTED'; return out
record('C14','identical ingest rerun','ALREADY_INGESTED',tC14)

# C15 conflicting ingest rerun.
def tC15():
 p=ROOT/'c15.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c15m'; r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); ev=bad_evidence(ROOT/'c15e.json','conflict_rerun')
 try:r.ingest_observed(p,'2026-09-20',ev,md); raise AssertionError('passed')
 except r.RunnerError as e: assert 'conflicting ingest rerun' in e.reason; return e.payload()
record('C15','conflicting ingest rerun','hard fail',tC15)

# C16 tamper observed/test/forecast manifests -> hash fail.
def tC16():
 # observed
 p=ROOT/'c16.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c16m'; md.mkdir(); r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); op=r._find_manifest_for_observed(md,'2026-09-20'); obj=json.loads(op.read_text()); obj['payload']['reason']='TAMPER'; op.write_text(json.dumps(obj))
 fails=[]
 try:r.load_workbook_state(p,md)
 except r.RunnerError as e:fails.append('observed')
 # test evidence
 te=make_test_evidence(ROOT/'c16te.json'); obj=json.loads(te.read_text()); obj['payload']['passed']=33; te.write_text(json.dumps(obj))
 try:r._validate_test_evidence(te,r.runner_sha256())
 except r.RunnerError as e:fails.append('test_evidence')
 # forecast manifest
 p2=ROOT/'c16f.xlsx'; shutil.copy2(BASE,p2); md2=ROOT/'c16fm'; md2.mkdir(); f=seed_forecast(p2,md2); obj=json.loads(f['manifest'].read_text()); obj['payload']['m0_uniform']=0.9; f['manifest'].write_text(json.dumps(obj))
 try:r.settle_forecast(p2,f['forecast_id'],'2026-09-20',REAL_EVIDENCE,md2)
 except r.RunnerError as e:fails.append('forecast')
 assert set(fails)=={'observed','test_evidence','forecast'}; return {'tamper_detected':sorted(fails)}
record('C16','manifest tampering','hash verification FAIL for observed/test/forecast',tC16)

# C17 concurrent RUN_FORECAST => one commit, one row.
def tC17():
 p=ROOT/'c17.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c17m'; md.mkdir(); r.ingest_observed(p,'2026-09-20',REAL_EVIDENCE,md); te=make_test_evidence(ROOT/'c17te.json'); helper=ROOT/'concurrent_helper.py'; helper_script(helper)
 cmd=[sys.executable,str(helper),str(RUNNER),'run',str(p),str(md),str(te),'NA','NA']
 a=subprocess.Popen(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True); b=subprocess.Popen(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True); ao,ae=a.communicate(); bo,be=b.communicate();
 wb=load_workbook(p); led=wb['PROSPECTIVE_LEDGER']; rows=[rr for rr in range(2,led.max_row+1) if led.cell(rr,1).value not in (None,'')]; wb.close(); successes=sum(x==0 for x in (a.returncode,b.returncode)); assert successes==1 and len(rows)==1; return {'returncodes':[a.returncode,b.returncode],'ledger_rows':len(rows)}
record('C17','concurrent RUN_FORECAST','only one commit and one Ledger row',tC17)

# C18 concurrent settlement => one consistent final state, no double append.
def tC18():
 p=ROOT/'c18.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c18m'; md.mkdir(); f=seed_forecast(p,md); helper=ROOT/'concurrent_helper2.py'; helper_script(helper)
 cmd=[sys.executable,str(helper),str(RUNNER),'settle',str(p),str(md),'NA',f['forecast_id'],str(REAL_EVIDENCE)]
 a=subprocess.Popen(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True); b=subprocess.Popen(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True); ao,ae=a.communicate(); bo,be=b.communicate();
 wb=load_workbook(p); led=wb['PROSPECTIVE_LEDGER']; rec=wb['RECONCILIATION_1200']; settled=[rr for rr in range(2,led.max_row+1) if str(led.cell(rr,16).value or '')=='SETTLED_VERIFIED']; rows=[rr for rr in range(2,rec.max_row+1) if rec.cell(rr,1).value not in (None,'')]; wb.close(); successes=sum(x==0 for x in (a.returncode,b.returncode)); assert successes==1 and len(settled)==1 and len(rows)==1201; return {'returncodes':[a.returncode,b.returncode],'settled_rows':len(settled),'draw_rows':len(rows)}
record('C18','concurrent settlement','exactly one final state and no double append',tC18)

# C19 external workbook change after staging -> compare-and-commit fail, external change preserved.
def tC19():
 p=ROOT/'c19.xlsx'; shutil.copy2(BASE,p); md=ROOT/'c19m'; md.mkdir(); st=r.load_workbook_state(p,MAN); wb=load_workbook(p); set_activation(wb,'c19_candidate','WOULD_COMMIT'); mp=md/'c19.json'; manifest=r._manifest_with_hash({'artifact_type':'TEST_TXN_MANIFEST_V1','marker':'C19'}); oldfault=r._fault
 def mutate(point):
  if point=='AFTER_STAGE':
   ext=load_workbook(p); set_activation(ext,'external_modification','EXTERNAL'); ext.save(p); ext.close()
 r._fault=mutate
 try:
  try:r.commit_workbook_and_manifest(p,wb,mp,manifest,st.workbook_sha256,'C19_TEST'); raise AssertionError('passed')
  except r.RunnerError as e: assert e.status.startswith('STATE CHANGED DURING RUN'); err=e.payload()
 finally:r._fault=oldfault; wb.close()
 assert not mp.exists(); ext=load_workbook(p); kv={str(ext['MODEL_ACTIVATION'].cell(rr,1).value or ''):ext['MODEL_ACTIVATION'].cell(rr,2).value for rr in range(2,ext['MODEL_ACTIVATION'].max_row+1)}; ext.close(); assert kv.get('external_modification')=='EXTERNAL' and 'c19_candidate' not in kv; return err
record('C19','workbook external modification / TOCTOU','abort, no manifest, no candidate commit',tC19)

passed=sum(x['status']=='PASS' for x in RESULTS); failed=len(RESULTS)-passed
report={'test_suite_revision':r.TEST_SUITE_REVISION,'runner_sha256':r.runner_sha256(),'spec_hash':r.SPEC_HASH,'master_content_sha256':r.FROZEN_MASTER_1200_SHA256,'software_test_fixture_notice':'Base workbook fixture clears stale mirror reconciliation flags ONLY for software testing. It is not production data certification.','executed_at':datetime.now(timezone.utc).isoformat(),'runtime':r.runtime_manifest(),'passed':passed,'failed':failed,'tests':RESULTS}
(OUT/'R4_34_TEST_RESULTS.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
# Only issue TEST_EVIDENCE_MANIFEST if 34/34 actual assertions passed.
if passed==34 and failed==0:
 payload={'artifact_type':'TEST_EVIDENCE_MANIFEST_V1','test_suite_revision':r.TEST_SUITE_REVISION,'spec_hash':r.SPEC_HASH,'runner_sha256':r.runner_sha256(),'master_content_sha256':r.FROZEN_MASTER_1200_SHA256,'runtime_observed':r.runtime_manifest(),'runtime_exact_production_match':(f"{sys.version_info.major}.{sys.version_info.minor}"==r.RUNTIME_LOCK['python_major_minor'] and all(r.runtime_manifest()[k]==r.RUNTIME_LOCK[k] for k in ('numpy','scipy','scikit_learn','openpyxl_operational'))),'runtime_note':'Authorization requires this flag TRUE; current host may remain non-production even when all software assertions pass.','passed':passed,'failed':failed,'executed_at':report['executed_at'],'tests':[{'id':x['id'],'status':x['status'],'input_hash_before':x['input_hash_before'],'input_hash_after':x['input_hash_after'],'manifest_hashes_after':x['manifest_hashes_after']} for x in RESULTS],'evidence_payload_sha256':r.sha256_bytes(r.canonical_json(RESULTS).encode())}
 obj=r._manifest_with_hash(payload); (OUT/'TEST_EVIDENCE_MANIFEST_V1.json').write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':passed,'failed':failed,'output':str(OUT)},indent=2))

#!/usr/bin/env python3
"""High-level operational wrapper for XSMB v2.1-LIVE-001.

User-facing operations:
  1) UPDATE_RESULT -> update-result
  2) RUN_NEXT      -> run-next

This wrapper does not change M7 mathematics. It delegates all model/data
integrity decisions to xsmb_reference_runner_R4_snapshot.py.
"""
from __future__ import annotations

import argparse
import json
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import xsmb_reference_runner_R4_snapshot as core

DEFAULT_WORKBOOK = "XSMB_RECONCILIATION_R5_1200_v2_1_PRODUCTION.xlsx"
DEFAULT_MANIFEST_DIR = "manifests"
DEFAULT_TEST_EVIDENCE = "test-evidence/TEST_EVIDENCE_MANIFEST_V1.json"
APPROVED_PREDECESSOR_RUNNER_SHA256 = "7c9f1bf773a484c20a09942a90a31613179963869dca30e65be8541d5bc7fd7b"
R4_CERTIFICATION_SHA256 = "93c4e1cea98518fa023cd1c31a71b0c495388cd132faaf82e91a48394500aabe"
DEFAULT_R4_CERTIFICATION = str(Path(__file__).resolve().parent / "certification" / "R4_SNAPSHOT_CERTIFICATION.json")


def _latest_and_rows(workbook: str, manifest_dir: str) -> tuple[core.WorkbookState, str]:
    state = core.load_workbook_state(workbook, manifest_dir)
    if not state.draws.dates:
        raise core.RunnerError("IMPLEMENTATION GATE FAIL — DO NOT FORECAST", "no verified draws in production workbook")
    return state, state.draws.dates[-1]


def update_result(
    workbook: str,
    actual_date: str,
    evidence_json: str,
    manifest_dir: str,
) -> dict[str, Any]:
    """Settle an existing forecast or ingest a missed draw without backfill.

    The wrapper resolves forecast_id automatically. It never fabricates a
    forecast row for a date that had no pre-draw forecast.
    """
    actual_date = core.normalize_date(actual_date)
    state, latest = _latest_and_rows(workbook, manifest_dir)

    rows = [
        r for r in state.ledger_records
        if r.get("forecast_date") and core.normalize_date(r["forecast_date"]) == actual_date
    ]
    if len(rows) > 1:
        raise core.RunnerError(
            "HASH/STATE CONFLICT — DO NOT FORECAST",
            "multiple forecast rows exist for actual_date",
            {"actual_date": actual_date, "count": len(rows)},
        )

    if len(rows) == 1:
        result = core.settle_forecast(
            workbook_path=workbook,
            forecast_id=str(rows[0]["forecast_id"]),
            actual_date=actual_date,
            evidence_json=evidence_json,
            manifest_dir=manifest_dir,
        )
        mode = "SETTLE_EXISTING_FORECAST"
    else:
        # No pre-draw row: only the explicit observed-no-forecast path is legal.
        result = core.ingest_observed(
            workbook_path=workbook,
            actual_date=actual_date,
            evidence_json=evidence_json,
            manifest_dir=manifest_dir,
            reason="MISSED_PRE_DRAW_CUTOFF",
        )
        mode = "INGEST_OBSERVED_NO_FORECAST"

    new_state, new_latest = _latest_and_rows(workbook, manifest_dir)
    next_date = (date.fromisoformat(new_latest) + timedelta(days=1)).isoformat()
    return {
        "operation": "UPDATE_RESULT",
        "mode": mode,
        "actual_date": actual_date,
        "previous_latest_verified_date": latest,
        "latest_verified_date": new_latest,
        "next_forecast_date": next_date,
        "research_state": new_state.activation.get("research_state"),
        "result": result,
    }


def _validate_r4_certification(certification_path: str) -> dict[str, Any]:
    path = Path(certification_path)
    if not path.exists():
        raise core.RunnerError("IMPLEMENTATION GATE FAIL — DO NOT FORECAST", "missing R4 snapshot certification")
    if core.sha256_file(path) != R4_CERTIFICATION_SHA256:
        raise core.RunnerError("IMPLEMENTATION GATE FAIL — DO NOT FORECAST", "R4 snapshot certification file hash mismatch")
    cert = json.loads(path.read_text(encoding="utf-8"))
    expected = {
        "status": "PASS",
        "operational_revision": "R4-SNAPSHOT",
        "model_spec": core.SPEC_CORE["spec_version"],
        "spec_hash": core.SPEC_HASH,
        "runner_sha256": core.runner_sha256(),
        "model_math_changed": False,
    }
    bad = {k: {"expected": v, "observed": cert.get(k)} for k, v in expected.items() if cert.get(k) != v}
    checks = cert.get("checks") or {}
    required_checks = (
        "legacy_hardening_34_of_34",
        "same_host_R3_R4_probability_identity",
        "snapshot_schema_and_hash",
        "snapshot_reconstructs_rank",
        "snapshot_reconstructs_probability_within_2e_12",
    )
    if bad or any(checks.get(k) is not True for k in required_checks):
        raise core.RunnerError(
            "IMPLEMENTATION GATE FAIL — DO NOT FORECAST",
            "R4 snapshot certification is not fully PASS",
            {"fields": bad, "checks": checks},
        )
    return cert


def _ensure_runner_activation(
    workbook: str,
    manifest_dir: str,
    test_evidence: str,
    certification_path: str,
) -> dict[str, Any]:
    """Promote R3 -> certified R4 only after all prior forecasts are settled.

    This changes only the operational runner pin in MODEL_ACTIVATION. It does
    not alter MODEL_SPEC, M7 mathematics, MASTER, P00-P99 or any prior forecast.
    """
    path = Path(workbook)
    mdir = Path(manifest_dir)
    current_sha = core.runner_sha256()

    # Validate R4 before touching the workbook.
    core.assert_runtime()
    test_manifest = core._validate_test_evidence(Path(test_evidence), current_sha)
    cert = _validate_r4_certification(certification_path)

    with core.workbook_lock(path, True):
        core.recover_transactions(path)
        state = core.load_workbook_state(path, mdir)
        active_sha = str(state.activation.get("reference_runner_sha256") or "")

        if active_sha == current_sha:
            return {
                "status": "ALREADY_ACTIVE",
                "runner_sha256": current_sha,
                "operational_revision": core.OPERATIONAL_REVISION,
            }

        if active_sha != APPROVED_PREDECESSOR_RUNNER_SHA256:
            raise core.RunnerError(
                "HASH MISMATCH — DO NOT FORECAST",
                "runner promotion predecessor is not the approved R3 runner",
                {"observed": active_sha, "approved_predecessor": APPROVED_PREDECESSOR_RUNNER_SHA256},
            )

        unsettled = [
            {
                "forecast_id": r.get("forecast_id"),
                "forecast_date": r.get("forecast_date"),
                "record_status": r.get("record_status"),
            }
            for r in state.ledger_records
            if str(r.get("record_status") or "") == "LOCKED_PRE_DRAW"
        ]
        if unsettled:
            raise core.RunnerError(
                "PRE-DRAW GATE FAIL — DO NOT FORECAST",
                "runner promotion blocked until all existing LOCKED_PRE_DRAW forecasts are settled",
                unsettled,
            )

        wb = load_workbook(path, read_only=False, data_only=False, keep_links=True)
        core._update_activation_field(wb, "reference_runner_sha256", current_sha)
        core._update_activation_field(wb, "operational_revision", core.OPERATIONAL_REVISION)
        core._update_activation_field(wb, "runner_promotion_status", "CERTIFIED_R4_SNAPSHOT")
        core._update_activation_field(
            wb,
            "runner_promotion_test_evidence_manifest_sha256",
            str(test_manifest.get("manifest_sha256") or ""),
        )
        core._update_activation_field(wb, "runner_promotion_certification_sha256", R4_CERTIFICATION_SHA256)

        payload = {
            "artifact_type": "XSMB_RUNNER_PROMOTION_V1",
            "from_runner_sha256": active_sha,
            "to_runner_sha256": current_sha,
            "operational_revision": core.OPERATIONAL_REVISION,
            "model_spec_version": core.SPEC_CORE["spec_version"],
            "spec_hash": core.SPEC_HASH,
            "master_content_sha256": core.FROZEN_MASTER_1200_SHA256,
            "model_math_changed": False,
            "test_evidence_manifest_sha256": str(test_manifest.get("manifest_sha256") or ""),
            "r4_certification_file_sha256": R4_CERTIFICATION_SHA256,
            "r4_certification_runner_sha256": cert["runner_sha256"],
            "promoted_at_local": datetime.now(core.VIETNAM_TZ).isoformat(),
        }
        manifest = core._manifest_with_hash(payload)
        mpath = mdir / "XSMB_RUNNER_PROMOTION_R3_TO_R4_SNAPSHOT.json"
        commit = core.commit_workbook_and_manifest(
            path,
            wb,
            mpath,
            manifest,
            state.workbook_sha256,
            "RUNNER_PROMOTION_R3_TO_R4_SNAPSHOT",
        )
        wb.close()
        return {
            "status": "PROMOTED",
            "from_runner_sha256": active_sha,
            "runner_sha256": current_sha,
            "operational_revision": core.OPERATIONAL_REVISION,
            "promotion_manifest": str(mpath),
            **commit,
        }


def run_next(
    workbook: str,
    manifest_dir: str,
    test_evidence: str,
    forecast_date: str | None = None,
    schedule_evidence: str | None = None,
    r4_certification: str = DEFAULT_R4_CERTIFICATION,
) -> dict[str, Any]:
    """Run the next prospective forecast from the latest verified draw.

    data_cutoff is derived automatically. If forecast_date is omitted, the next
    calendar day is used. Core preflight remains fail-closed.
    """
    state, cutoff = _latest_and_rows(workbook, manifest_dir)
    target = core.normalize_date(forecast_date) if forecast_date else (
        date.fromisoformat(cutoff) + timedelta(days=1)
    ).isoformat()

    promotion = _ensure_runner_activation(
        workbook=workbook,
        manifest_dir=manifest_dir,
        test_evidence=test_evidence,
        certification_path=r4_certification,
    )
    state, cutoff = _latest_and_rows(workbook, manifest_dir)

    result = core.run_forecast(
        workbook_path=workbook,
        forecast_date=target,
        data_cutoff_s=cutoff,
        manifest_dir=manifest_dir,
        test_evidence=test_evidence,
        schedule_evidence=schedule_evidence,
    )
    return {
        "operation": "RUN_NEXT",
        "forecast_date": target,
        "data_cutoff": cutoff,
        "research_state": state.activation.get("research_state"),
        "runner_activation": promotion,
        "result": result,
    }


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="XSMB two-command operational wrapper")
    sub = p.add_subparsers(dest="command", required=True)

    u = sub.add_parser("update-result")
    u.add_argument("--actual-date", required=True)
    u.add_argument("--evidence-json", required=True)
    u.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    u.add_argument("--manifest-dir", default=DEFAULT_MANIFEST_DIR)

    r = sub.add_parser("run-next")
    r.add_argument("--forecast-date")
    r.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    r.add_argument("--manifest-dir", default=DEFAULT_MANIFEST_DIR)
    r.add_argument("--test-evidence", default=DEFAULT_TEST_EVIDENCE)
    r.add_argument("--schedule-evidence")
    r.add_argument("--r4-certification", default=DEFAULT_R4_CERTIFICATION)

    return p


def main() -> None:
    args = build_parser().parse_args()
    try:
        if args.command == "update-result":
            result = update_result(args.workbook, args.actual_date, args.evidence_json, args.manifest_dir)
        elif args.command == "run-next":
            result = run_next(
                args.workbook, args.manifest_dir, args.test_evidence,
                args.forecast_date, args.schedule_evidence, args.r4_certification,
            )
        else:
            raise core.RunnerError("IMPLEMENTATION GATE FAIL — DO NOT FORECAST", "unknown ops command")
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except core.RunnerError as exc:
        print(json.dumps(exc.payload(), ensure_ascii=False, indent=2))
        raise SystemExit(2) from exc


if __name__ == "__main__":
    main()

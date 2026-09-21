#!/usr/bin/env python3
from __future__ import annotations
import argparse, importlib.util, json, shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

GOVERNANCE_KEYS=[
 "proposed_version","status","activation_gate","spec_gate","research_state","null_control",
 "prospective_challenger","forecast_objective","live_checkpoints","spec_hash","data_basis",
 "prompt_status","heuristic_status","penalty_fix","pre_draw_lock","post_draw_lock",
 "no_target_leakage","no_backfill","freeze_basis","rule"
]
OPERATIONAL_KEYS=[
 "reference_runner_sha256","operational_revision","runner_promotion_status",
 "runner_promotion_test_evidence_manifest_sha256","runner_promotion_certification_sha256"
]

def load_runner(path: Path):
    spec=importlib.util.spec_from_file_location("xsmb_schema_runner",path)
    if spec is None or spec.loader is None: raise RuntimeError(f"cannot load runner {path}")
    mod=importlib.util.module_from_spec(spec); spec.loader.exec_module(mod); return mod

def activation_map(wb)->dict[str,Any]:
    sh=wb["MODEL_ACTIVATION"]; out={}
    for r in range(2,sh.max_row+1):
        k=str(sh.cell(r,1).value or "").strip()
        if k: out[k]=sh.cell(r,2).value
    return out

def norm(v:Any)->Any:
    if isinstance(v,(datetime,date)): return v.isoformat()
    if isinstance(v,(str,int,float,bool)) or v is None: return v
    return str(v)

def existing_value_hash(core:Any, wb, excluded:set[str])->str:
    payload=[]
    for name in wb.sheetnames:
        if name in excluded: continue
        sh=wb[name]
        rows=[]
        for row in sh.iter_rows():
            rows.append([norm(c.value) for c in row])
        payload.append({"sheet":name,"rows":rows})
    return core.sha256_bytes(core.canonical_json(payload).encode("utf-8"))

def write_kv_sheet(wb,name:str,rows:list[tuple[Any,Any,Any,Any]]):
    if name in wb.sheetnames: del wb[name]
    sh=wb.create_sheet(name)
    sh.append(["Field","Value","Status","Notes"])
    for row in rows: sh.append(list(row))

def migrate(input_path:Path, output_path:Path, manifest_dir:Path, runner_path:Path, promotion_manifest:Path|None)->dict[str,Any]:
    core=load_runner(runner_path)
    before_state=core.load_workbook_state(input_path,manifest_dir)
    wb=load_workbook(input_path,read_only=False,data_only=False,keep_links=True)
    legacy_hash_before=existing_value_hash(core,wb,{"MODEL_GOVERNANCE","OPERATIONAL_STATE","RUNNER_HISTORY"})
    a=activation_map(wb)

    gov=[]
    for k in GOVERNANCE_KEYS:
        gov.append((k,a.get(k),"ACTIVE" if k in a else "MISSING","Migrated from MODEL_ACTIVATION; model-governance field only."))
    gov.extend([
        ("model_governance_schema_version","XSMB_MODEL_GOVERNANCE_V1","ACTIVE","Operational fields are intentionally excluded."),
        ("model_math_changed",False,"LOCKED","OPS hardening must not change M7 mathematics."),
    ])
    write_kv_sheet(wb,"MODEL_GOVERNANCE",gov)

    ops=[
        ("operational_state_schema_version","XSMB_OPERATIONAL_STATE_V1","ACTIVE","Sole operational-state sheet for future hardened runner."),
        ("active_runner_sha256",a.get("reference_runner_sha256"),"ACTIVE","Migrated from current live runner pin."),
        ("operational_revision",a.get("operational_revision") or "R3-HARDENED","ACTIVE","Current operational runner revision."),
        ("active_test_evidence_manifest_sha256",a.get("runner_promotion_test_evidence_manifest_sha256") or a.get("test_evidence_manifest_sha256"),"ACTIVE","Current exact-runtime authorization evidence."),
        ("active_certification_sha256",a.get("runner_promotion_certification_sha256"),"ACTIVE","Current runner certification when applicable."),
        ("research_state",a.get("research_state") or "NO VERIFIED EDGE","ACTIVE","Descriptive research state; not a forecast override."),
        ("legacy_model_activation_policy","READ_ONLY_ARCHIVE","LOCKED","MODEL_ACTIVATION is retained for audit compatibility but is not operational authority after promotion."),
    ]
    write_kv_sheet(wb,"OPERATIONAL_STATE",ops)

    if "RUNNER_HISTORY" in wb.sheetnames: del wb["RUNNER_HISTORY"]
    hs=wb.create_sheet("RUNNER_HISTORY")
    hs.append(["Event ID","Event Type","From Runner SHA256","To Runner SHA256","Operational Revision","Test Evidence SHA256","Certification SHA256","Model Math Changed","Event Time Local","Source Manifest SHA256"])
    active=str(a.get("reference_runner_sha256") or "")
    if promotion_manifest:
        pm=core._validate_payload_manifest(promotion_manifest,"XSMB_RUNNER_PROMOTION_V1")
        p=pm["payload"]
        if str(p.get("to_runner_sha256") or "")!=active:
            raise core.RunnerError("OPS SCHEMA FAIL — DO NOT PROMOTE","promotion manifest does not terminate at current active runner",{"active":active,"manifest_to":p.get("to_runner_sha256")})
        hs.append([
            f"PROMOTION-{str(pm.get('manifest_sha256') or '')[:16]}","RUNNER_PROMOTION",
            p.get("from_runner_sha256"),p.get("to_runner_sha256"),p.get("operational_revision"),
            p.get("test_evidence_manifest_sha256"),p.get("r4_certification_file_sha256"),
            p.get("model_math_changed"),p.get("promoted_at_local"),pm.get("manifest_sha256")
        ])
    else:
        hs.append(["BOOTSTRAP-CURRENT","BOOTSTRAP_ACTIVE_RUNNER",None,active,a.get("operational_revision") or "R3-HARDENED",a.get("test_evidence_manifest_sha256"),None,False,None,None])

    output_path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(output_path); wb.close()

    check=load_workbook(output_path,read_only=False,data_only=False,keep_links=True)
    legacy_hash_after=existing_value_hash(core,check,{"MODEL_GOVERNANCE","OPERATIONAL_STATE","RUNNER_HISTORY"})
    check.close()
    if legacy_hash_after!=legacy_hash_before:
        raise core.RunnerError("OPS SCHEMA FAIL — DO NOT PROMOTE","legacy workbook values changed during schema migration",{"before":legacy_hash_before,"after":legacy_hash_after})

    after_state=core.load_workbook_state(output_path,manifest_dir)
    if before_state.draws.dates!=after_state.draws.dates or before_state.draws.lotto!=after_state.draws.lotto or before_state.ledger_records!=after_state.ledger_records:
        raise core.RunnerError("OPS SCHEMA FAIL — DO NOT PROMOTE","draws or Ledger changed during schema migration")
    if core.canonical_master_hash(after_state.draws)!=core.FROZEN_MASTER_1200_SHA256:
        raise core.RunnerError("OPS SCHEMA FAIL — DO NOT PROMOTE","frozen MASTER hash changed")

    return {
        "status":"PASS",
        "schema_revision":"OPS_HARDENING_V1",
        "model_math_changed":False,
        "legacy_values_hash_before":legacy_hash_before,
        "legacy_values_hash_after":legacy_hash_after,
        "input_workbook_sha256":before_state.workbook_sha256,
        "output_workbook_sha256":core.sha256_file(output_path),
        "frozen_master_sha256":core.FROZEN_MASTER_1200_SHA256,
        "ledger_records":len(after_state.ledger_records),
        "latest_verified_draw":after_state.draws.dates[-1],
        "new_sheets":["MODEL_GOVERNANCE","OPERATIONAL_STATE","RUNNER_HISTORY"]
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--input",required=True)
    p.add_argument("--output",required=True)
    p.add_argument("--manifest-dir",required=True)
    p.add_argument("--runner",default="xsmb_runtime/xsmb_reference_runner_R4_snapshot.py")
    p.add_argument("--promotion-manifest")
    p.add_argument("--report",required=True)
    a=p.parse_args()
    report=migrate(Path(a.input),Path(a.output),Path(a.manifest_dir),Path(a.runner),Path(a.promotion_manifest) if a.promotion_manifest else None)
    Path(a.report).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=="__main__":
    main()
